"""Document generation: presentations, spreadsheets, Word files and PDFs.

Ports ``actions/office_builder.py``, ``actions/docx_tools.py`` and
``actions/pdf_tools.py``. The generation libraries (python-pptx, python-docx,
openpyxl, reportlab) are all cross-platform, so only two things changed:

  * ``os.startfile(path)`` became ``xdg-open`` through the apps adapter.
  * Output defaults to the XDG documents directory instead of the Desktop.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from .. import config
from ..platform import apps, shell
from ..registry import skill
from ..safety import Danger, check_path, enforce


def _output_path(filename: str, extension: str) -> Path:
    """Resolve where a generated document should land."""
    name = (filename or "").strip() or f"brahma-{datetime.now():%Y%m%d-%H%M%S}"
    path = config.resolve_user_path(name)

    # A bare name goes to the documents directory, not the CWD.
    if not path.is_absolute() or str(path) == name:
        path = config.user_dirs()["documents"] / name

    if path.suffix.lower() != extension:
        path = path.with_suffix(extension)
    return path


def _finish(path: Path, kind: str, open_after: bool) -> str:
    size = path.stat().st_size if path.exists() else 0
    message = f"{kind} saved: {path} ({size / 1024:.0f} KB)"
    if open_after:
        message += "\n" + apps.open_uri(str(path))
    return message


@skill(category="documents", danger=Danger.MODERATE)
def create_presentation(
    title: str,
    slides: list[dict],
    filename: str = "",
    open_after: bool = False,
    _source: str = "local",
    _confirm: bool = False,
) -> str:
    """Build a PowerPoint deck.

    Args:
        title: Deck title, used for the opening slide.
        slides: One dict per slide with 'title' and 'bullets' (a list of strings).
        filename: Output name. Defaults to the documents folder.
        open_after: Open the file when it is done.
    """
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError:
        return "python-pptx is not installed. Run: pip install python-pptx"

    path = _output_path(filename or title, ".pptx")
    enforce(check_path(path, write=True), f"write presentation {path}",
            skill="create_presentation", source=_source, pre_confirmed=_confirm)

    deck = Presentation()

    # Title slide
    title_slide = deck.slides.add_slide(deck.slide_layouts[0])
    title_slide.shapes.title.text = title
    if len(title_slide.placeholders) > 1:
        title_slide.placeholders[1].text = f"Generated {datetime.now():%d %B %Y}"

    for entry in slides or []:
        if not isinstance(entry, dict):
            continue
        layout = deck.slide_layouts[1]
        slide = deck.slides.add_slide(layout)
        slide.shapes.title.text = str(entry.get("title", "Slide"))

        bullets = entry.get("bullets") or entry.get("points") or []
        if isinstance(bullets, str):
            bullets = [b.strip() for b in bullets.split("\n") if b.strip()]

        body = slide.placeholders[1].text_frame
        body.clear()
        for i, bullet in enumerate(bullets):
            para = body.paragraphs[0] if i == 0 else body.add_paragraph()
            para.text = str(bullet)
            para.level = 0
            para.font.size = Pt(18)

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        deck.save(str(path))
    except OSError as exc:
        return f"Could not save the presentation: {exc}"

    return _finish(path, f"Presentation ({len(deck.slides)} slides)", open_after)


@skill(category="documents", danger=Danger.MODERATE)
def create_spreadsheet(
    sheets: list[dict],
    filename: str = "",
    open_after: bool = False,
    _source: str = "local",
    _confirm: bool = False,
) -> str:
    """Build an Excel workbook.

    Args:
        sheets: One dict per sheet with 'name', 'headers' (list) and 'rows'
            (list of lists).
        filename: Output name. Defaults to the documents folder.
        open_after: Open the file when it is done.
    """
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return "openpyxl is not installed. Run: pip install openpyxl"

    path = _output_path(filename, ".xlsx")
    enforce(check_path(path, write=True), f"write spreadsheet {path}",
            skill="create_spreadsheet", source=_source, pre_confirmed=_confirm)

    book = Workbook()
    book.remove(book.active)

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", start_color="2F5496")

    for entry in sheets or []:
        if not isinstance(entry, dict):
            continue
        sheet = book.create_sheet(str(entry.get("name", "Sheet"))[:31])

        headers = entry.get("headers") or []
        if headers:
            sheet.append([str(h) for h in headers])
            for cell in sheet[1]:
                cell.font = header_font
                cell.fill = header_fill

        for row in entry.get("rows") or []:
            sheet.append(list(row) if isinstance(row, (list, tuple)) else [row])

        # Size columns to their content so the file is readable on open.
        for column in range(1, sheet.max_column + 1):
            width = max(
                (len(str(sheet.cell(row=r, column=column).value or ""))
                 for r in range(1, min(sheet.max_row, 200) + 1)),
                default=10,
            )
            sheet.column_dimensions[get_column_letter(column)].width = min(max(width + 2, 10), 60)

    if not book.sheetnames:
        book.create_sheet("Sheet1")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        book.save(str(path))
    except OSError as exc:
        return f"Could not save the workbook: {exc}"

    return _finish(path, f"Workbook ({len(book.sheetnames)} sheets)", open_after)


@skill(category="documents", danger=Danger.MODERATE)
def create_document(
    title: str,
    sections: list[dict],
    filename: str = "",
    open_after: bool = False,
    _source: str = "local",
    _confirm: bool = False,
) -> str:
    """Build a Word document.

    Args:
        title: Document title.
        sections: One dict per section with 'heading' and 'body' text.
        filename: Output name. Defaults to the documents folder.
        open_after: Open the file when it is done.
    """
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError:
        return "python-docx is not installed. Run: pip install python-docx"

    path = _output_path(filename or title, ".docx")
    enforce(check_path(path, write=True), f"write document {path}",
            skill="create_document", source=_source, pre_confirmed=_confirm)

    doc = Document()
    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    stamp = doc.add_paragraph(f"Generated {datetime.now():%d %B %Y}")
    stamp.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for entry in sections or []:
        if not isinstance(entry, dict):
            continue
        if entry.get("heading"):
            doc.add_heading(str(entry["heading"]), level=int(entry.get("level", 1)))
        body = entry.get("body") or entry.get("text") or ""
        for paragraph in str(body).split("\n\n"):
            if paragraph.strip():
                doc.add_paragraph(paragraph.strip())
        for bullet in entry.get("bullets") or []:
            doc.add_paragraph(str(bullet), style="List Bullet")

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(path))
    except OSError as exc:
        return f"Could not save the document: {exc}"

    return _finish(path, "Document", open_after)


@skill(category="documents", danger=Danger.MODERATE)
def create_pdf(
    title: str,
    content: str,
    filename: str = "",
    open_after: bool = False,
    _source: str = "local",
    _confirm: bool = False,
) -> str:
    """Build a PDF from a title and body text.

    Args:
        title: Document title.
        content: Body text. Blank lines separate paragraphs; a line ending
            in ':' is treated as a heading.
        filename: Output name. Defaults to the documents folder.
        open_after: Open the file when it is done.
    """
    try:
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import cm
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError:
        return "reportlab is not installed. Run: pip install reportlab"

    path = _output_path(filename or title, ".pdf")
    enforce(check_path(path, write=True), f"write PDF {path}",
            skill="create_pdf", source=_source, pre_confirmed=_confirm)

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        "BrahmaTitle", parent=styles["Title"], alignment=TA_CENTER, spaceAfter=18
    )

    story = [Paragraph(_escape(title), title_style), Spacer(1, 0.4 * cm)]

    for block in str(content).split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.endswith(":") and len(block) < 90:
            story.append(Paragraph(_escape(block.rstrip(":")), styles["Heading2"]))
        else:
            story.append(Paragraph(_escape(block), styles["BodyText"]))
        story.append(Spacer(1, 0.25 * cm))

    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        SimpleDocTemplate(
            str(path), pagesize=A4,
            leftMargin=2 * cm, rightMargin=2 * cm,
            topMargin=2 * cm, bottomMargin=2 * cm,
            title=title,
        ).build(story)
    except Exception as exc:
        return f"Could not build the PDF: {exc}"

    return _finish(path, "PDF", open_after)


def _escape(text: str) -> str:
    """reportlab's Paragraph parses a mini-HTML dialect."""
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace("\n", "<br/>")
    )


@skill(category="documents", danger=Danger.MODERATE)
def convert_document(
    source: str,
    target_format: str = "pdf",
    _source: str = "local",
    _confirm: bool = False,
) -> str:
    """Convert a document to another format using LibreOffice.

    Args:
        source: File to convert.
        target_format: Target extension, e.g. 'pdf', 'docx', 'odt', 'csv'.
    """
    binary = shell.first_available("libreoffice", "soffice")
    if not binary:
        return (
            "LibreOffice is not installed, so conversion is unavailable.\n"
            "Install it with:  sudo pacman -S libreoffice-fresh"
        )

    src = config.resolve_user_path(source)
    if not src.exists():
        return f"Not found: {src}"

    out_dir = src.parent
    enforce(check_path(out_dir / f"{src.stem}.{target_format}", write=True),
            f"convert {src.name} to {target_format}",
            skill="convert_document", source=_source, pre_confirmed=_confirm)

    result = shell.run(
        [binary, "--headless", "--convert-to", target_format,
         "--outdir", str(out_dir), str(src)],
        timeout=180,
    )
    if not result.ok:
        return f"Conversion failed: {result.output}"

    produced = out_dir / f"{src.stem}.{target_format}"
    if produced.exists():
        return f"Converted: {produced}"
    return f"LibreOffice reported success but {produced} is missing.\n{result.stdout}"
